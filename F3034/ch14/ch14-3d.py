from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.units import pixels_to_EMU

wb = load_workbook("月產品銷售含直播.xlsx")
ws = wb.active
 
data = Reference(ws,
                 min_col = 2, 
                 max_col = 4,
                 min_row = 1,
                 max_row = 7)
  
chart = LineChart()
chart.add_data(data,
               titles_from_data=True)
dates = Reference(ws, min_col=1, min_row=2, max_row=7)
chart.set_categories(dates)
chart.title = "各通路產品銷售圖表"
chart.style = 6
chart.x_axis.title = "月份"
chart.y_axis.title = "銷售額"
chart.legend.position = "b"

s1 = chart.series[0]
s1.marker.symbol = "triangle"
s1.marker.graphicalProperties.solidFill = "FF0000" 
s1.marker.graphicalProperties.line.solidFill = "FF0000"
s1.graphicalProperties.line.noFill = True

s2 = chart.series[1]
s2.graphicalProperties.line.solidFill = "00AAAA"
s2.graphicalProperties.line.dashStyle = "sysDot"
s2.graphicalProperties.line.width = pixels_to_EMU(2.5)

s3 = chart.series[2]
s3.smooth = True # 建立平滑線

data2 = Reference(ws,
                  min_col = 5,
                  min_row = 1,
                  max_row = 7)  
chart2 = LineChart()
chart2.add_data(data2,
                titles_from_data=True)
chart2.set_categories(dates)
chart2.y_axis.title = "直播銷售額"
chart2.y_axis.axId = 200
chart2.y_axis.crosses = "max"
chart += chart2

ws.add_chart(chart, "G2")
wb.save("月產品銷售含直播圖表.xlsx")
wb.close()