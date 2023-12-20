from openpyxl import load_workbook
from openpyxl.chart import AreaChart3D,Reference

wb = load_workbook("產品銷售.xlsx")
ws = wb.active
 
data = Reference(ws, min_col = 2, max_col = 4,
                 min_row = 1, max_row = 7)
  
chart = AreaChart3D()
chart.add_data(data, titles_from_data=True)
chart.title = "產品銷售3D面積圖"
chart.style = 15
chart.x_axis.title = "商品"
chart.y_axis.title = "銷售"
ws.add_chart(chart, "F2")
wb.save("areaChart3D.xlsx")
wb.close()