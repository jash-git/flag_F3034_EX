from openpyxl import load_workbook
from openpyxl.chart import BarChart3D, Reference

wb = load_workbook("產品銷售.xlsx")
ws = wb.active
 
data = Reference(ws, min_col = 2, max_col = 4,
                 min_row = 1, max_row = 7)

chart = BarChart3D()
chart.add_data(data, titles_from_data=True)
chart.title = "產品銷售3D長條圖"
chart.style = 6
chart.x_axis.title = "商品"
chart.y_axis.title = "銷售"
ws.add_chart(chart, "F2")
wb.save("barChart3D.xlsx")
wb.close()