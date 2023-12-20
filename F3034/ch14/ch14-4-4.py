from openpyxl import load_workbook
from openpyxl.chart import StockChart,Reference

wb = load_workbook("2330TW.xlsx")
ws = wb.active
 
data = Reference(ws, min_col = 2, max_col = 5,
                 min_row = 1, max_row = 19)
  
chart = StockChart()
chart.add_data(data, from_rows=False,
               titles_from_data=True)
dates = Reference(ws, min_col=1, min_row=2, max_row=19)
chart.set_categories(dates)
chart.title = "台積電股票圖"
chart.style = 20
chart.x_axis.title = "日期"
chart.y_axis.title = "股價"
ws.add_chart(chart, "H2")
wb.save("stockChart.xlsx")
wb.close()