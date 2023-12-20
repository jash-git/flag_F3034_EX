from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series
 
wb = load_workbook("球員薪水.xlsx")
ws = wb.active

chart = ScatterChart()
ydata = Reference(ws, min_col = 3, min_row = 2, max_row = 98)
xdata = Reference(ws, min_col = 5, min_row = 2, max_row = 98)

series = Series(values = ydata, xvalues = xdata, title = "2018")
series.marker.symbol = "circle"
series.marker.graphicalProperties.solidFill = "FF0000" 
series.marker.graphicalProperties.line.solidFill = "FF0000"
series.graphicalProperties.line.noFill = True
chart.series.append(series)
chart.title = "NBA球員薪水與得分散佈圖"
chart.x_axis.title = "得分"
chart.y_axis.title = "薪水"
ws.add_chart(chart, "H2")
wb.save("scatterChart.xlsx")
wb.close()