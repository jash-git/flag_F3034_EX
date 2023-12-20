from openpyxl import load_workbook
from openpyxl.chart import LineChart,Reference

wb = load_workbook("2330TW.xlsx")
ws = wb.active
 
data = Reference(ws, min_col = 2, max_col = 5,
                 min_row = 1, max_row = 19)
  
chart = LineChart()
chart.add_data(data, from_rows=False,
               titles_from_data=True)
chart.title = "股價折線圖"
chart.style = 10
chart.x_axis.title = "日期"
chart.y_axis.title = "股價"
ws.add_chart(chart, "H2")
wb.save("lineChart.xlsx")
wb.close()