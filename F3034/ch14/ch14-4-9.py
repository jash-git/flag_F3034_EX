from openpyxl import load_workbook
from openpyxl.chart import RadarChart, Reference
 
wb = load_workbook("植物類月銷售.xlsx")
ws = wb.active

labels = Reference(ws, min_col = 1, min_row = 2,
                   max_row = 13)
data = Reference(ws, min_col = 2, max_col = 5, 
                 min_row = 2, max_row = 13)

chart = RadarChart()
chart.type = "filled" 
chart.add_data(data, titles_from_data = True)
chart.set_categories(labels)
chart.title = "植物類商品月銷售雷達圖"
chart.style = 26
chart.y_axis.delete = True
ws.add_chart(chart, "G2")
wb.save("radarChart.xlsx")
wb.close()