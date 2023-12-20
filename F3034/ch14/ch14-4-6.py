from openpyxl import load_workbook
from openpyxl.chart import DoughnutChart, Reference
 
wb = load_workbook("口味銷售量.xlsx")
ws = wb.active

labels = Reference(ws, min_col = 1, min_row = 2, max_row = 5)
data = Reference(ws, min_col = 2, min_row = 1, max_row = 5)
 
chart = DoughnutChart()
chart.add_data(data, titles_from_data = True)
chart.set_categories(labels)
chart.title = "口味銷售量的甜甜圈圖"
ws.add_chart(chart, "D2")
wb.save("doughnutChart.xlsx")
wb.close()
