from openpyxl import load_workbook
from openpyxl.chart import AreaChart,Reference

wb = load_workbook("產品銷售.xlsx")
ws = wb.active
 
data = Reference(ws, min_col = 2, max_col = 2,
                 min_row = 1, max_row = 7)
  
chart = AreaChart()
chart.add_data(data, titles_from_data=True)
chart.title = "網路商店銷售面積圖"
chart.x_axis.title = "商品"
chart.y_axis.title = "銷售"
ws.add_chart(chart, "F2")
wb.save("areaChart.xlsx")
wb.close()