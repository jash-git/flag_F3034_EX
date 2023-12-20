from openpyxl import load_workbook
from openpyxl.chart import BubbleChart, Series, Reference

wb = load_workbook("產品銷售2.xlsx")
ws = wb.active

chart = BubbleChart()
ydata = Reference(ws, min_col = 2, min_row = 2, max_row = 5)
xdata = Reference(ws, min_col = 3, min_row = 2, max_row = 5)
size = Reference(ws, min_col = 1, min_row = 2, max_row = 5)

series = Series(values = ydata, xvalues = xdata,
                zvalues = size, title = "2018")
chart.series.append(series)
chart.title = "2018年網路與實體銷售泡泡圖"
chart.x_axis.title = "實體店面"
chart.y_axis.title = "網路商店"
ws.add_chart(chart, "E2")
wb.save("bubbleChart.xlsx")
wb.close()