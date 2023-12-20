from openpyxl import load_workbook
from openpyxl.chart import PieChart3D, Reference
 
wb = load_workbook("口味銷售量.xlsx")
ws = wb.active

labels = Reference(ws, min_col = 1, min_row = 2, max_row = 5)
data = Reference(ws, min_col = 2, min_row = 1, max_row = 5)

chart = PieChart3D()
chart.add_data(data, titles_from_data = True)
chart.set_categories(labels)
chart.title = "口味銷售量3D派圖"
ws.add_chart(chart, "D2")
wb.save("pieChart3D.xlsx")
wb.close()

