from openpyxl import load_workbook
from openpyxl.chart import SurfaceChart3D, Reference
 
wb = load_workbook("曲面圖測試資料.xlsx")
ws = wb.active

labels = Reference(ws, min_col = 1, min_row = 2,
                   max_row = 10)
data = Reference(ws, min_col = 2, max_col = 6, 
                 min_row = 1, max_row = 10) 

chart = SurfaceChart3D()
chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)
chart.title = "3D曲面圖"
chart.style = 26
ws.add_chart(chart, "H2")
wb.save("surfaceChart3D.xlsx")
wb.close()