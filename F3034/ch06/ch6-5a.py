from openpyxl import load_workbook

wb = load_workbook("營業額和股票表格.xlsx")
ws = wb.active

print("表格數:", len(ws.tables))
# 切換表格
tab = ws.tables["Table1"] 
print(tab.name)
# 走訪工作表的表格
for table in ws.tables.values():
    print(table.name)
# 取得表格資訊串列
print(ws.tables.items())
# 刪除表格
del ws.tables["Table1"]
print("表格數:", len(ws.tables))
