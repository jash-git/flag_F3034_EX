from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = load_workbook("營業額.xlsx")
ws = wb.active

tab = Table(displayName="Table1", 
            ref="A1:D7")
style = TableStyleInfo(name="TableStyleMedium10", 
                       showFirstColumn=False,
                       showLastColumn=False,
                       showRowStripes=True,
                       showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)
wb.save("營業額_Table.xlsx") 
wb.close()
